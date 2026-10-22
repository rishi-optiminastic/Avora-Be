"""The monthly attendance register as a full, honest month.

Pins the fixes to the monthly report and its .xlsx:

- a new joiner is never marked absent for days before they joined;
- a leaver still appears for the days they worked in the month (HR settles
  their final pay from it), and is not marked absent after they left;
- the Daily sheet has a row for every calendar day, weekends and holidays
  labelled, so it really is a register;
- an unreal month (2026-13) is a 422, not a 500;
- the monthly late allowance counts lates the same way whether a day falls
  before or inside the requested range.

The month is a fixed past one (September 2025) so weekends are known: the 1st
is a Monday and the 7th, 14th, 21st and 28th are Sundays.
"""

from __future__ import annotations

from datetime import UTC, date, datetime
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.core.exceptions import ValidationError
from app.models.employee import Employee, EmployeeStatus
from app.models.holiday import Holiday
from app.models.resignation import Resignation, ResignationStatus
from app.models.work_session import WorkSession
from app.services.attendance_service import AttendanceService
from tests.conftest import _Seed, auth_headers

_MONTH = "2025-09"
_DAYS_IN_MONTH = 30
_OFFICE = ZoneInfo("Asia/Kolkata")


def _at(day: date, hour: int, minute: int = 0) -> datetime:
    """An office wall-clock time, as the UTC instant actually stored."""
    return datetime(day.year, day.month, day.day, hour, minute, tzinfo=_OFFICE).astimezone(UTC)


def _sep(day: int) -> date:
    return date(2025, 9, day)


async def _session(
    db: AsyncSession,
    employee: Employee,
    day: date,
    start: tuple[int, int] = (9, 0),
    end: tuple[int, int] = (18, 0),
) -> None:
    db.add(
        WorkSession(
            employee_id=employee.id,
            clock_in_at=_at(day, *start),
            clock_out_at=_at(day, *end),
            source="dashboard",
        )
    )
    await db.commit()


async def _deactivate(db: AsyncSession, employee: Employee) -> None:
    employee.is_active = False
    employee.status = EmployeeStatus.INACTIVE
    await db.commit()


async def _report(
    client: AsyncClient, settings: Settings, who: Employee, month: str = _MONTH
) -> dict[str, dict[str, object]]:
    res = await client.get(
        f"/api/v1/attendance/report?month={month}", headers=auth_headers(settings, who)
    )
    assert res.status_code == 200, res.text
    return {row["employee_id"]: row for row in res.json()}


async def _daily_sheet(
    client: AsyncClient, settings: Settings, who: Employee
) -> list[dict[str, object]]:
    res = await client.get(
        f"/api/v1/attendance/report/export?month={_MONTH}", headers=auth_headers(settings, who)
    )
    assert res.status_code == 200, res.text
    wb = load_workbook(BytesIO(res.content))
    sheet = wb[wb.sheetnames[1]]
    headers = [c.value for c in sheet[1]]
    return [
        dict(zip(headers, [c.value for c in row], strict=True))
        for row in sheet.iter_rows(min_row=2)
    ]


async def _summary_sheet(
    client: AsyncClient, settings: Settings, who: Employee
) -> list[dict[str, object]]:
    """The first sheet of the register - one row per employee for the month."""
    res = await client.get(
        f"/api/v1/attendance/report/export?month={_MONTH}", headers=auth_headers(settings, who)
    )
    assert res.status_code == 200, res.text
    wb = load_workbook(BytesIO(res.content))
    sheet = wb[wb.sheetnames[0]]
    headers = [c.value for c in sheet[1]]
    return [
        dict(zip(headers, [c.value for c in row], strict=True))
        for row in sheet.iter_rows(min_row=2)
    ]


def _rows_for(sheet: list[dict[str, object]], person: Employee) -> dict[str, dict[str, object]]:
    return {str(r["Date"]): r for r in sheet if r["Employee"] == person.full_name}


async def _absent_days_in_range(
    client: AsyncClient, settings: Settings, seed: _Seed, start: date, end: date
) -> int:
    """How many working days in [start, end] a person with no sessions is absent
    for - the holiday/weekend-aware yardstick, taken from the manager."""
    res = await client.get(
        f"/api/v1/attendance/range?start={start}&end={end}",
        headers=auth_headers(settings, seed.admin),
    )
    assert res.status_code == 200, res.text
    return sum(
        1
        for r in res.json()
        if r["employee_id"] == str(seed.manager.id) and r["status"] == "absent"
    )


# --------------------------------------------------------------------------- #
# New joiners
# --------------------------------------------------------------------------- #
async def test_days_before_joining_are_not_absences(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    seed.report.hire_date = _sep(15)
    await db.commit()

    report = await _report(client, settings, seed.admin)
    expected = await _absent_days_in_range(client, settings, seed, _sep(15), _sep(30))

    assert report[str(seed.report.id)]["absent_days"] == expected
    assert report[str(seed.manager.id)]["absent_days"] > expected


async def test_daily_sheet_marks_days_before_joining(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    seed.report.hire_date = _sep(15)
    await db.commit()

    mine = _rows_for(await _daily_sheet(client, settings, seed.admin), seed.report)

    assert all(mine[_sep(d).isoformat()]["Status"] == "Not yet joined" for d in range(1, 15))
    assert mine[_sep(15).isoformat()]["Status"] == "Absent"  # a Monday, no session


async def test_someone_hired_after_the_month_is_not_listed(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    seed.report.hire_date = date(2025, 10, 1)
    await db.commit()

    report = await _report(client, settings, seed.admin)

    assert str(seed.report.id) not in report


async def test_range_view_skips_pre_hire_days_too(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """The employment window lives in the shared absence synthesis, so the
    date-range view stops inventing pre-joining absences as well."""
    seed.report.hire_date = _sep(15)
    await db.commit()

    res = await client.get(
        f"/api/v1/attendance/range?start={_sep(1)}&end={_sep(30)}",
        headers=auth_headers(settings, seed.admin),
    )

    assert res.status_code == 200
    days = {r["day"] for r in res.json() if r["employee_id"] == str(seed.report.id)}
    assert days
    assert min(days) >= _sep(15).isoformat()


async def test_single_day_view_skips_a_day_before_joining(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """The single-day view uses the same employment window as the range and
    month views: no invented Absent before someone joined, the joining day
    itself counts, and a real punch before it is still shown."""
    seed.report.hire_date = _sep(15)
    await db.commit()

    async def _day(day: date) -> dict[str, str]:
        res = await client.get(
            f"/api/v1/attendance?date={day}", headers=auth_headers(settings, seed.admin)
        )
        assert res.status_code == 200, res.text
        return {r["employee_id"]: r["status"] for r in res.json()}

    before = await _day(_sep(10))  # a Wednesday
    assert str(seed.report.id) not in before
    assert before[str(seed.manager.id)] == "absent"
    assert (await _day(_sep(15)))[str(seed.report.id)] == "absent"  # a Monday

    await _session(db, seed.report, _sep(11))
    assert str(seed.report.id) in await _day(_sep(11))


# --------------------------------------------------------------------------- #
# Leavers
# --------------------------------------------------------------------------- #
async def test_leaver_deactivated_mid_month_is_still_in_the_report(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    await _session(db, seed.report, _sep(3))
    await _deactivate(db, seed.report)

    report = await _report(client, settings, seed.admin)

    row = report[str(seed.report.id)]
    assert row["full_days"] == 1
    # No resignation on file: the last recorded day stands in for the leaving
    # date, so the days after it are not absences.
    expected = await _absent_days_in_range(client, settings, seed, _sep(1), _sep(2))
    assert row["absent_days"] == expected


async def test_report_rows_carry_name_and_department_for_leavers(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """The live directory drops leavers, so the report row itself names them -
    otherwise the on-screen report and its CSV read "Unknown"."""
    seed.report.department = "Design"
    await _session(db, seed.report, _sep(3))
    await _deactivate(db, seed.report)

    report = await _report(client, settings, seed.admin)

    row = report[str(seed.report.id)]
    assert row["employee_name"] == seed.report.full_name
    assert row["department"] == "Design"
    assert report[str(seed.manager.id)]["employee_name"] == seed.manager.full_name


async def test_leaver_daily_sheet_shows_worked_days_and_left_after(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    await _session(db, seed.report, _sep(3))
    db.add(
        Resignation(
            employee_id=seed.report.id,
            last_working_day=_sep(10),
            status=ResignationStatus.ACCEPTED,
        )
    )
    await db.commit()
    await _deactivate(db, seed.report)

    mine = _rows_for(await _daily_sheet(client, settings, seed.admin), seed.report)

    assert mine[_sep(3).isoformat()]["Check in"] == "09:00"
    # The accepted resignation's last working day bounds employment, so the
    # unattended working days up to it are absences and nothing after is.
    assert mine[_sep(10).isoformat()]["Status"] == "Absent"
    assert all(mine[_sep(d).isoformat()]["Status"] == "Left" for d in range(11, 31))


async def test_leaver_with_no_trace_in_the_month_is_not_listed(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    await _session(db, seed.report, date(2025, 8, 4))
    await _deactivate(db, seed.report)

    report = await _report(client, settings, seed.admin)

    assert str(seed.report.id) not in report


async def test_leaver_is_scoped_like_everyone_else(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """Including leavers widens the time window, never the people: the manager
    still sees their (departed) report, an outsider still does not."""
    await _session(db, seed.report, _sep(3))
    await _deactivate(db, seed.report)

    as_manager = await _report(client, settings, seed.manager)
    as_outsider = await _report(client, settings, seed.outsider)

    assert str(seed.report.id) in as_manager
    assert str(seed.report.id) not in as_outsider
    assert set(as_outsider) == {str(seed.outsider.id)}


async def test_leavers_do_not_leak_into_the_live_roster(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """`all_in_scope` (every other feature) still excludes inactive people."""
    await _session(db, seed.report, _sep(3))
    await _deactivate(db, seed.report)

    res = await client.get(
        "/api/v1/attendance/range?start=2025-09-03&end=2025-09-03",
        headers=auth_headers(settings, seed.admin),
    )

    assert res.status_code == 200
    assert str(seed.report.id) not in {r["employee_id"] for r in res.json()}


# --------------------------------------------------------------------------- #
# Every calendar day
# --------------------------------------------------------------------------- #
async def test_daily_sheet_has_every_day_with_weekends_and_holidays(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    db.add(Holiday(name="Founders Day", date=_sep(2)))
    await db.commit()
    await _session(db, seed.report, _sep(7), (10, 0), (14, 0))  # worked a Sunday

    sheet = await _daily_sheet(client, settings, seed.admin)
    mine = _rows_for(sheet, seed.report)
    theirs = _rows_for(sheet, seed.manager)

    assert len(mine) == _DAYS_IN_MONTH
    assert len(theirs) == _DAYS_IN_MONTH
    assert theirs[_sep(2).isoformat()]["Status"] == "Holiday: Founders Day"
    assert theirs[_sep(14).isoformat()]["Status"] == "Weekend"
    sunday = mine[_sep(7).isoformat()]
    assert str(sunday["Status"]).startswith("Weekend (")
    assert sunday["Check in"] == "10:00"
    assert sunday["Check out"] == "14:00"
    assert sunday["Day"] == "Sun"


async def test_a_weekend_shift_counts_as_a_present_day(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """Present days are the days you were actually at work, not the working
    days you were at work. Someone who came in on a Sunday is present that
    Sunday, and the Summary says so.

    Not a new rule - `_summarize` counts statuses and never filtered by
    weekday - but nothing pinned it, so a future "only count working days"
    tidy-up would have silently erased weekend work from everyone's month.
    """
    await _session(db, seed.report, _sep(1))  # Monday
    await _session(db, seed.report, _sep(7), (10, 0), (18, 0))  # Sunday

    res = await client.get(
        f"/api/v1/attendance/report?month={_MONTH}", headers=auth_headers(settings, seed.admin)
    )

    assert res.status_code == 200, res.text
    mine = next(r for r in res.json() if r["employee_id"] == str(seed.report.id))
    assert mine["full_days"] == 2  # the Monday AND the Sunday

    summary = await _summary_sheet(client, settings, seed.admin)
    row = next(r for r in summary if r["Employee"] == seed.report.full_name)
    assert row["Present days"] == 2


# --------------------------------------------------------------------------- #
# Invalid month
# --------------------------------------------------------------------------- #
@pytest.mark.parametrize("month", ["2026-13", "2026-00", "2026-1", "0000-01", "9999-12"])
@pytest.mark.parametrize("path", ["/api/v1/attendance/report", "/api/v1/attendance/report/export"])
async def test_an_unreal_month_is_a_422(
    client: AsyncClient, settings: Settings, seed: _Seed, path: str, month: str
) -> None:
    res = await client.get(f"{path}?month={month}", headers=auth_headers(settings, seed.admin))

    assert res.status_code == 422
    assert "YYYY-MM" in res.json()["error"]["message"]


def test_the_service_also_rejects_an_unreal_month() -> None:
    """Defence in depth: a non-route caller gets a 422-mapped error, not a 500."""
    with pytest.raises(ValidationError):
        AttendanceService._month_bounds("2026-13")
    assert AttendanceService._month_bounds("2024-02") == ("2024-02-01", "2024-02-29")


# --------------------------------------------------------------------------- #
# Late counting
# --------------------------------------------------------------------------- #
async def test_lates_before_a_range_count_like_lates_inside_it(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """Three manual (non-biometric) lates, then a fourth. The fourth is a half
    day whether the range starts on the 1st or on the 4th itself - the lates
    before a range used to count biometric punches only, so the same day read
    as a full day in one view and a half day in another."""
    for day in (1, 2, 3, 4):
        await _session(db, seed.report, _sep(day), (10, 30), (19, 0))

    async def status_on_4th(start: date) -> str:
        res = await client.get(
            f"/api/v1/attendance/range?start={start}&end={_sep(4)}",
            headers=auth_headers(settings, seed.admin),
        )
        assert res.status_code == 200, res.text
        rows = [
            r
            for r in res.json()
            if r["employee_id"] == str(seed.report.id) and r["day"] == _sep(4).isoformat()
        ]
        assert len(rows) == 1
        return str(rows[0]["status"])

    assert await status_on_4th(_sep(1)) == "half_day"
    assert await status_on_4th(_sep(4)) == "half_day"

    report = await _report(client, settings, seed.admin)
    assert report[str(seed.report.id)]["late_days"] == 0  # all met 8h from arrival
    assert report[str(seed.report.id)]["full_days"] == 3
    assert report[str(seed.report.id)]["half_days"] == 1
