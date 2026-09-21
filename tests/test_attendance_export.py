"""The monthly attendance register (.xlsx).

Two things matter beyond "does a file come back". The sheet must agree with the
API report - they read the same rollup, so a month can never be counted twice
two ways - and the export must not widen scope: an export that shows more than
the screen is a data leak with a filename (CLAUDE.md §9, security rule 5.3).
"""

from __future__ import annotations

from datetime import UTC, date, datetime, timedelta
from io import BytesIO

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import Settings
from app.models.work_session import WorkSession
from tests.conftest import _Seed, auth_headers

_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _recent_weekday() -> date:
    today = datetime.now(UTC).date()
    return today - timedelta(days=today.weekday())


def _month_of(day: date) -> str:
    return f"{day.year:04d}-{day.month:02d}"


async def _worked(db: AsyncSession, employee_id: object, day: date) -> None:
    """A 9:30-to-18:00 day, in UTC (the test org runs on UTC)."""
    db.add(
        WorkSession(
            employee_id=employee_id,
            clock_in_at=datetime(day.year, day.month, day.day, 9, 30, tzinfo=UTC),
            clock_out_at=datetime(day.year, day.month, day.day, 18, 0, tzinfo=UTC),
            source="dashboard",
        )
    )
    await db.commit()


async def _export(client: AsyncClient, settings: Settings, who: object, month: str):
    return await client.get(
        f"/api/v1/attendance/report/export?month={month}",
        headers=auth_headers(settings, who),  # type: ignore[arg-type]
    )


def _sheets(content: bytes):
    return load_workbook(BytesIO(content))


async def test_export_returns_a_two_sheet_workbook(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    day = _recent_weekday()
    await _worked(db, seed.report.id, day)

    res = await _export(client, settings, seed.admin, _month_of(day))

    assert res.status_code == 200
    assert res.headers["content-type"].startswith(_XLSX)
    assert f"attendance-{_month_of(day)}.xlsx" in res.headers["content-disposition"]
    wb = _sheets(res.content)
    assert len(wb.sheetnames) == 2
    assert wb.sheetnames[0].startswith("Summary")
    assert wb.sheetnames[1].startswith("Daily")


async def test_daily_sheet_carries_check_in_and_check_out(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """The whole point of the report: the actual in/out times, not just totals."""
    day = _recent_weekday()
    await _worked(db, seed.report.id, day)

    res = await _export(client, settings, seed.admin, _month_of(day))

    daily = _sheets(res.content)[_sheets(res.content).sheetnames[1]]
    headers = [c.value for c in daily[1]]
    assert "Check in" in headers and "Check out" in headers
    rows = [
        dict(zip(headers, [c.value for c in row], strict=True))
        for row in daily.iter_rows(min_row=2)
    ]
    mine = [
        r for r in rows if r["Employee"] == seed.report.full_name and r["Date"] == day.isoformat()
    ]
    assert len(mine) == 1
    assert mine[0]["Check in"] == "09:30"
    assert mine[0]["Check out"] == "18:00"
    assert mine[0]["Hours"] == 8.5


async def test_summary_sheet_agrees_with_the_api_report(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """Both read the same rollup, so the sheet can never contradict the screen."""
    day = _recent_weekday()
    await _worked(db, seed.report.id, day)
    month = _month_of(day)

    api = await client.get(
        f"/api/v1/attendance/report?month={month}", headers=auth_headers(settings, seed.admin)
    )
    res = await _export(client, settings, seed.admin, month)

    assert api.status_code == 200
    by_id = {r["employee_id"]: r for r in api.json()}
    wb = _sheets(res.content)
    summary = wb[wb.sheetnames[0]]
    headers = [c.value for c in summary[1]]
    rows = {
        r[headers.index("Employee")]: r
        for r in ([c.value for c in row] for row in summary.iter_rows(min_row=2))
    }
    mine = rows[seed.report.full_name]
    expected = by_id[str(seed.report.id)]
    assert mine[headers.index("Full days")] == expected["full_days"]
    assert mine[headers.index("Absent days")] == expected["absent_days"]
    assert mine[headers.index("Total hours")] == round(expected["worked_minutes"] / 60, 2)


async def test_export_is_scoped_like_every_other_attendance_read(
    client: AsyncClient, db: AsyncSession, settings: Settings, seed: _Seed
) -> None:
    """An employee's export contains only themselves. Exporting must never be a
    way to see people the screen would not show."""
    day = _recent_weekday()
    await _worked(db, seed.report.id, day)
    await _worked(db, seed.outsider.id, day)

    res = await _export(client, settings, seed.outsider, _month_of(day))

    assert res.status_code == 200
    wb = _sheets(res.content)
    summary = wb[wb.sheetnames[0]]
    names = {row[0].value for row in summary.iter_rows(min_row=2)}
    assert names == {seed.outsider.full_name}
    assert seed.report.full_name not in names


async def test_a_bad_month_is_rejected(
    client: AsyncClient, settings: Settings, seed: _Seed
) -> None:
    res = await _export(client, settings, seed.admin, "not-a-month")

    assert res.status_code == 422


async def test_unauthenticated_is_rejected(client: AsyncClient) -> None:
    res = await client.get("/api/v1/attendance/report/export?month=2026-09")

    assert res.status_code == 401
