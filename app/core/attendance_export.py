"""The monthly attendance register as .xlsx.

Two sheets, because the month is two questions. **Summary** answers "how did
each person's month go" - one row per employee, the day counts and total hours.
**Daily** answers "what happened on the 14th" - one row per employee per day,
with the check-in and check-out times.

Times are rendered in the org's timezone, already converted by the caller. A UTC
timestamp in a sheet HR reads is a support ticket waiting to happen.
"""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from typing import cast

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill("solid", fgColor="5A48E0")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_MAX_WIDTH = 38


@dataclass(frozen=True, slots=True)
class AttendanceSummaryRow:
    employee_name: str
    department: str
    full_days: int
    half_days: int
    late_days: int
    absent_days: int
    leave_days: int
    regularized_days: int
    present_days: int
    worked_hours: float

    def as_cells(self) -> list[object]:
        return [
            self.employee_name,
            self.department,
            self.present_days,
            self.full_days,
            self.half_days,
            self.late_days,
            self.absent_days,
            self.leave_days,
            self.regularized_days,
            self.worked_hours,
        ]


@dataclass(frozen=True, slots=True)
class AttendanceDailyRow:
    employee_name: str
    department: str
    day: str
    weekday: str
    status: str
    check_in: str
    check_out: str
    worked_hours: float
    late: str
    regularized: str
    in_source: str
    out_source: str

    def as_cells(self) -> list[object]:
        return [
            self.employee_name,
            self.department,
            self.day,
            self.weekday,
            self.status,
            self.check_in,
            self.check_out,
            self.worked_hours,
            self.late,
            self.regularized,
            self.in_source,
            self.out_source,
        ]


_SUMMARY_HEADERS = [
    "Employee",
    "Department",
    "Present days",
    "Full days",
    "Half days",
    "Late days",
    "Absent days",
    "Leave days",
    "Regularised days",
    "Total hours",
]
_DAILY_HEADERS = [
    "Employee",
    "Department",
    "Date",
    "Day",
    "Status",
    "Check in",
    "Check out",
    "Hours",
    "Late",
    "Regularised",
    "In source",
    "Out source",
]
# Columns holding hours, so they read 7.50 rather than 7.5 or 7.499999.
_SUMMARY_HOURS = {10}
_DAILY_HOURS = {8}


def _write(ws: Worksheet, headers: list[str], rows: list[list[object]], hours: set[int]) -> None:
    ws.append(headers)
    for col in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in rows:
        ws.append(row)
        for col in hours:
            ws.cell(row=ws.max_row, column=col).number_format = "0.00"
    ws.freeze_panes = "A2"
    for col, header in enumerate(headers, start=1):
        widest = max(
            [len(header)] + [len(str(r[col - 1])) for r in rows if r[col - 1] is not None] or [0]
        )
        ws.column_dimensions[get_column_letter(col)].width = min(widest + 3, _MAX_WIDTH)


def build_attendance_xlsx(
    summary: list[AttendanceSummaryRow],
    daily: list[AttendanceDailyRow],
    *,
    month_label: str,
) -> bytes:
    """Render the month to .xlsx bytes: a Summary sheet and a Daily sheet."""
    wb = Workbook()
    first = cast(Worksheet, wb.active)  # a fresh Workbook always has an active sheet
    first.title = f"Summary {month_label}"[:31]  # Excel caps a sheet name at 31 chars
    _write(first, _SUMMARY_HEADERS, [r.as_cells() for r in summary], _SUMMARY_HOURS)

    second = wb.create_sheet(f"Daily {month_label}"[:31])
    _write(second, _DAILY_HEADERS, [r.as_cells() for r in daily], _DAILY_HOURS)

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
